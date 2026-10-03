import sys
from unittest.mock import MagicMock

# Preemptively stub onnxruntime before pymupdf4llm import to prevent native C-extension segfault
for _mod in ("onnxruntime", "onnxruntime.capi", "onnxruntime.capi._pybind_state"):
    if _mod not in sys.modules:
        sys.modules[_mod] = MagicMock()

import pytest
from pathlib import Path
from qmd.converters import convert_to_markdown

def test_convert_xlsx_chartsheet(tmp_path, monkeypatch):
    try:
        import openpyxl
    except ImportError:
        pytest.skip("openpyxl is not installed")

    from qmd.converters import convert_to_markdown

    class MockR:
        def __init__(self, t):
            self.t = t
    class MockP:
        def __init__(self, r_list):
            self.r = r_list
    class MockRich:
        def __init__(self, p_list):
            self.p = p_list
    class MockTx:
        def __init__(self, rich):
            self.rich = rich
    class MockTitleRich:
        def __init__(self, text):
            self.tx = MockTx(MockRich([MockP([MockR(text)])]))

    class MockChartSimple:
        def __init__(self, title_str):
            self.title = title_str

    class MockChartRich:
        def __init__(self, text):
            self.title = MockTitleRich(text)

    class MockChartsheet:
        # Crucially, no iter_rows attribute
        def __init__(self):
            self.charts = [
                MockChartRich("Sales Trend (Rich)"),
                MockChartSimple("Revenue (Simple)")
            ]

    class MockWorksheet:
        def iter_rows(self, values_only=True):
            yield ["Col1", "Col2"]
            yield [10, 20]

    class MockWorkbook:
        def __init__(self, *args, **kwargs):
            self.sheetnames = ["Data", "Charts"]
        def __getitem__(self, item):
            if item == "Data":
                return MockWorksheet()
            return MockChartsheet()
        def close(self):
            pass

    monkeypatch.setattr(openpyxl, "load_workbook", MockWorkbook)

    xlsx_file = tmp_path / "test.xlsx"
    xlsx_file.write_bytes(b"dummy xlsx")

    md = convert_to_markdown(xlsx_file)
    
    # Verify standard sheet parsed
    assert "## Sheet: Data" in md
    assert "| Col1 | Col2 |" in md
    assert "| 10 | 20 |" in md
    
    # Verify chart sheet parsed correctly
    assert "## Sheet: Charts" in md
    assert "*[Chart Sheet]*" in md
    assert "- Chart Title: Sales Trend (Rich)" in md
    assert "- Chart Title: Revenue (Simple)" in md

def test_convert_pptx_no_embedded_image(tmp_path, monkeypatch):
    import pptx
    from qmd.converters import convert_to_markdown

    class MockShape:
        has_table = False
        has_text_frame = True

        @property
        def image(self):
            raise ValueError("no embedded image")

        @property
        def text_frame(self):
            class Paragraph:
                text = "Slide content paragraph"
                level = 0
            class TextFrame:
                text = "Slide content paragraph"
                paragraphs = [Paragraph()]
            return TextFrame()

    class MockShapes(list):
        title = None

    class MockSlide:
        shapes = MockShapes([MockShape()])

    class MockPresentation:
        slides = [MockSlide()]

    monkeypatch.setattr(pptx, "Presentation", lambda path: MockPresentation())

    dummy_pptx = tmp_path / "sample.pptx"
    dummy_pptx.write_bytes(b"dummy pptx")

    md = convert_to_markdown(dummy_pptx)
    assert "## Slide 1" in md
    assert "Slide content paragraph" in md

def test_convert_pptx_no_embedded_image_with_vision_config(tmp_path, monkeypatch):
    import pptx
    from qmd.converters import convert_to_markdown
    from qmd.config import Config

    class MockShape:
        has_table = False
        has_text_frame = True

        @property
        def image(self):
            raise ValueError("no embedded image")

        @property
        def text_frame(self):
            class Paragraph:
                text = "Slide text with vision config"
                level = 0
            class TextFrame:
                text = "Slide text with vision config"
                paragraphs = [Paragraph()]
            return TextFrame()

    class MockShapes(list):
        title = None

    class MockSlide:
        shapes = MockShapes([MockShape()])

    class MockPresentation:
        slides = [MockSlide()]

    monkeypatch.setattr(pptx, "Presentation", lambda path: MockPresentation())

    dummy_pptx = tmp_path / "sample_vision.pptx"
    dummy_pptx.write_bytes(b"dummy pptx")

    cfg = Config.from_dict({"vision_url": "http://127.0.0.1:9999"})
    errors = []
    md = convert_to_markdown(dummy_pptx, config=cfg, errors_out=errors)
    assert "## Slide 1" in md
    assert "Slide text with vision config" in md
    assert len(errors) == 0

def test_convert_pptx_real_presentation(tmp_path):
    try:
        from pptx import Presentation
    except ImportError:
        pytest.skip("python-pptx is not installed")

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Intro Title"
    slide.placeholders[1].text = "Subtitle Body"

    pptx_path = tmp_path / "real_test.pptx"
    prs.save(str(pptx_path))

    md = convert_to_markdown(pptx_path)
    assert "## Slide 1: Intro Title" in md
    assert "Subtitle Body" in md