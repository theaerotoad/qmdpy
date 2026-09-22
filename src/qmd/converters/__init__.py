"""
Document and Text Converters for QMD.

Converts non-markdown formats (.docx, .pptx, .xlsx, .csv, .html, .epub, .pdf) into Markdown text.
Organized as a modular package with math, image vision, and PDF processing submodules.
"""
import csv
import io
import re
import sys
from pathlib import Path
from typing import Union, List, Optional, Dict, Any

from .math import (
    _mathml_to_latex,
    _omml_to_latex,
    _extract_text_and_math,
)
from .images import (
    _clean_vision_markdown,
    clean_vision_text,
    _wrap_vision_xml,
    _is_image_processing_enabled,
    _is_verbose,
    _process_image_multimodal_llm,
    _process_image,
    _process_images_concurrently,
    _process_image_vision_api,
)
from .pdf import (
    _convert_pdf,
    build_header_detector,
    clean_heading_text,
    normalize_markdown_headings,
)


def _get_dplib():
    try:
        import dplib
        return dplib
    except ImportError:
        cur = Path(__file__).resolve().parent
        while cur.name != "src" and cur.parent != cur:
            cur = cur.parent
        src_dir = str(cur) if cur.name == "src" else str(Path(__file__).resolve().parents[2])
        if src_dir not in sys.path:
            sys.path.insert(0, src_dir)
        try:
            import dplib
            return dplib
        except ImportError:
            return None


SUPPORTED_EXTENSIONS = {
    ".md", ".markdown", ".txt",
    ".pdf",
    ".docx",
    ".pptx",
    ".xlsx",
    ".csv",
    ".html", ".htm",
    ".epub",
    ".mobi"
}


def is_supported_file(file_path: Union[str, Path]) -> bool:
    ext = Path(file_path).suffix.lower()
    return ext in SUPPORTED_EXTENSIONS


def _sanitize_text(text: str) -> str:
    """
    Sanitizes string by replacing invalid unicode surrogates (U+D800 - U+DFFF)
    with replacement characters so it can be encoded to UTF-8 without errors.
    """
    if not text:
        return ""
    try:
        return text.encode('utf-16', 'surrogatepass').decode('utf-16', 'replace')
    except Exception:
        return text


def _format_matrix_to_md_table(matrix: List[List[str]]) -> str:
    if not matrix:
        return ""

    max_cols = max(len(row) for row in matrix)
    if max_cols == 0:
        return ""

    norm_matrix = []
    for row in matrix:
        norm_row = row + [""] * (max_cols - len(row))
        norm_row = [re.sub(r'\s+', ' ', cell).strip().replace("|", "\\|") for cell in norm_row]
        norm_matrix.append(norm_row)

    col_is_empty = [True] * max_cols
    for row in norm_matrix:
        for c in range(max_cols):
            if row[c] != "":
                col_is_empty[c] = False

    if all(col_is_empty):
        return ""

    cols_to_keep = []
    c = 0
    while c < max_cols:
        if col_is_empty[c]:
            j = c
            while j < max_cols and col_is_empty[j]:
                j += 1
            empty_count = j - c
            if empty_count >= 3:
                cols_to_keep.append((False, empty_count))
                c = j
                continue
        cols_to_keep.append((True, c))
        c += 1

    row_is_empty = [all(cell == "" for cell in row) for row in norm_matrix]
    
    final_matrix = []
    i = 0
    while i < len(norm_matrix):
        if row_is_empty[i]:
            j = i
            while j < len(norm_matrix) and row_is_empty[j]:
                j += 1
            empty_count = j - i
            if empty_count >= 3:
                new_row = []
                marker_placed = False
                for keep, val in cols_to_keep:
                    if not keep:
                        if i == 0:
                            new_row.append(f"<{val} empty cols>")
                        else:
                            new_row.append("...")
                    else:
                        if not marker_placed:
                            new_row.append(f"<{empty_count} empty rows skipped>")
                            marker_placed = True
                        else:
                            new_row.append("")
                final_matrix.append(new_row)
                i = j
                continue
            
        row = norm_matrix[i]
        new_row = []
        for keep, val in cols_to_keep:
            if not keep:
                if i == 0:
                    new_row.append(f"<{val} empty cols>")
                else:
                    new_row.append("...")
            else:
                new_row.append(row[val])
        final_matrix.append(new_row)
        i += 1

    if not final_matrix:
        return ""

    header = final_matrix[0]
    separator = ["---"] * len(header)
    rows = final_matrix[1:] if len(final_matrix) > 1 else []

    header_line = "| " + " | ".join(header) + " |"
    sep_line = "| " + " | ".join(separator) + " |"
    body_lines = ["| " + " | ".join(row) + " |" for row in rows]

    return "\n".join([header_line, sep_line] + body_lines)


def _convert_text(path: Path) -> str:
    raw_bytes = path.read_bytes()

    if b'\x00' in raw_bytes[:8192]:
        raise ValueError(f"File '{path.name}' appears to be a binary file and cannot be converted as plain text.")

    for enc in ("utf-8", "utf-8-sig"):
        try:
            return raw_bytes.decode(enc)
        except UnicodeDecodeError:
            continue

    try:
        return raw_bytes.decode("latin-1")
    except Exception as e:
        raise ValueError(f"Could not decode text file '{path.name}': {e}")


def _convert_docx(path: Path, config=None, errors_out: Optional[List[dict]] = None) -> str:
    try:
        import docx
    except ImportError:
        raise ImportError("python-docx is required for converting .docx files. Install with `pip install python-docx`.")

    doc = docx.Document(str(path))
    md_lines: List[str] = []

    for elem in doc.element.body:
        tag = elem.tag.split('}')[-1]
        if tag == 'p':
            p = docx.text.paragraph.Paragraph(elem, doc)
            
            if config and _is_image_processing_enabled(config):
                p_images = []
                for blip in elem.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}blip'):
                    embed_id = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                    if embed_id and embed_id in doc.part.related_parts:
                        image_part = doc.part.related_parts[embed_id]
                        image_bytes = image_part.blob
                        filename = getattr(image_part, "filename", "image.png")
                        p_images.append((image_bytes, filename))
                if p_images:
                    results = _process_images_concurrently(p_images, config, errors_out=errors_out)
                    for vision_md in results:
                        if vision_md:
                            md_lines.append(vision_md + "\n")
            
            text = _extract_text_and_math(elem).strip()
            if not text:
                continue
            style_name = p.style.name.lower() if p.style else ""
            if "heading 1" in style_name:
                md_lines.append(f"# {text}\n")
            elif "heading 2" in style_name:
                md_lines.append(f"## {text}\n")
            elif "heading 3" in style_name:
                md_lines.append(f"### {text}\n")
            elif "heading 4" in style_name:
                md_lines.append(f"#### {text}\n")
            elif "heading" in style_name:
                md_lines.append(f"##### {text}\n")
            elif "list" in style_name or "bullet" in style_name:
                md_lines.append(f"- {text}")
            else:
                md_lines.append(f"{text}\n")

        elif tag == 'tbl':
            tbl = docx.table.Table(elem, doc)
            table_md = _format_matrix_to_md_table([
                [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                for row in tbl.rows
            ])
            if table_md:
                md_lines.append(table_md + "\n")

    return "\n".join(md_lines).strip()


def _convert_pptx(path: Path, config=None, errors_out: Optional[List[dict]] = None) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        raise ImportError("python-pptx is required for converting .pptx files. Install with `pip install python-pptx`.")

    prs = Presentation(str(path))
    md_lines: List[str] = []

    for i, slide in enumerate(prs.slides, 1):
        slide_title = ""
        slide_texts = []
        slide_images = []

        for shape in slide.shapes:
            shape_image = None
            if config and _is_image_processing_enabled(config):
                try:
                    shape_image = getattr(shape, "image", None)
                except Exception:
                    shape_image = None

            image_extracted = False
            if shape_image is not None:
                try:
                    image_bytes = shape_image.blob
                    filename = getattr(shape_image, "filename", "slide_image.png")
                    slide_images.append((image_bytes, filename))
                    image_extracted = True
                except Exception:
                    pass

            if not image_extracted and getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text.strip()
                if not text:
                    continue
                if shape == getattr(slide.shapes, "title", None):
                    if hasattr(shape, "element"):
                        parsed_title = _extract_text_and_math(shape.element).strip()
                        slide_title = parsed_title.replace('\n', ' ')
                    else:
                        slide_title = text.replace('\n', ' ')
                else:
                    for paragraph in shape.text_frame.paragraphs:
                        if hasattr(paragraph, "_p"):
                            p_text = _extract_text_and_math(paragraph._p).strip()
                        else:
                            p_text = paragraph.text.strip()
                        if p_text:
                            level = getattr(paragraph, "level", 0)
                            indent = "  " * level
                            slide_texts.append(f"{indent}- {p_text}")
            elif not image_extracted and getattr(shape, "has_table", False):
                table_matrix = []
                for row in shape.table.rows:
                    table_matrix.append([cell.text.strip().replace('\n', ' ') for cell in row.cells])
                table_md = _format_matrix_to_md_table(table_matrix)
                if table_md:
                    slide_texts.append(table_md)

        if slide_images:
            results = _process_images_concurrently(slide_images, config, errors_out=errors_out)
            for vision_md in results:
                if vision_md:
                    slide_texts.append(vision_md)

        header = f"## Slide {i}"
        if slide_title:
            header += f": {slide_title}"
        md_lines.append(header + "\n")

        if slide_texts:
            md_lines.extend(slide_texts)
            md_lines.append("")

    return "\n".join(md_lines).strip()


def _convert_xlsx(path: Path, config=None, errors_out: Optional[List[dict]] = None) -> str:
    try:
        import openpyxl
    except ImportError:
        raise ImportError("openpyxl is required for converting .xlsx files. Install with `pip install openpyxl`.")

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    md_lines: List[str] = []

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        md_lines.append(f"## Sheet: {sheet_name}\n")

        if config and _is_image_processing_enabled(config):
            sheet_images = []
            for img in getattr(sheet, "_images", []):
                try:
                    if hasattr(img, 'ref') and hasattr(img.ref, 'getvalue'):
                        image_bytes = img.ref.getvalue()
                    elif hasattr(img, '_data') and callable(img._data):
                        image_bytes = img._data()
                    else:
                        continue
                    filename = "excel_image.png"
                    sheet_images.append((image_bytes, filename))
                except Exception:
                    pass
            if sheet_images:
                results = _process_images_concurrently(sheet_images, config, errors_out=errors_out)
                for vision_md in results:
                    if vision_md:
                        md_lines.append(vision_md + "\n")

        if not hasattr(sheet, "iter_rows"):
            md_lines.append("*[Chart Sheet]*\n")
            try:
                charts = getattr(sheet, "charts", getattr(sheet, "_charts", []))
                for chart in charts:
                    if hasattr(chart, "title") and chart.title is not None:
                        title_text = ""
                        if hasattr(chart.title, "tx") and hasattr(chart.title.tx, "rich") and chart.title.tx.rich:
                            for p in getattr(chart.title.tx.rich, "p", []):
                                for r in getattr(p, "r", []):
                                    if hasattr(r, "t"):
                                        title_text += str(r.t)
                        else:
                            title_text = str(chart.title)
                        
                        if title_text and title_text != "None":
                            md_lines.append(f"- Chart Title: {title_text.strip()}\n")
            except Exception:
                pass
            continue

        matrix = []
        for row in sheet.iter_rows(values_only=True):
            if not row or all(v is None for v in row):
                continue
            row_vals = [str(v).strip().replace('\n', ' ') if v is not None else "" for v in row]
            matrix.append(row_vals)

        table_md = _format_matrix_to_md_table(matrix)
        if table_md:
            md_lines.append(table_md + "\n")

    wb.close()
    return "\n".join(md_lines).strip()


def _convert_csv(path: Path) -> str:
    try:
        content = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        content = path.read_text(encoding="latin-1", errors="replace")

    reader = csv.reader(io.StringIO(content))
    matrix = []
    for row in reader:
        if row and any(cell.strip() for cell in row):
            matrix.append([cell.strip().replace('\n', ' ') for cell in row])

    return _format_matrix_to_md_table(matrix)


def _convert_epub(path: Path) -> str:
    from qmd.epub import convert_epub_to_markdown
    try:
        return convert_epub_to_markdown(path)
    except Exception as e:
        raise ValueError(f"Failed to parse EPUB file '{path.name}': {e}")


def _convert_mobi(path: Path) -> str:
    try:
        import mobi
    except ImportError:
        raise ImportError("mobi is required for converting .mobi files. Install with `pip install mobi`.")
    
    import shutil
    import re
    
    try:
        tempdir, filepath = mobi.extract(str(path))
    except Exception as e:
        raise ValueError(f"Failed to extract MOBI file '{path.name}': {e}")
        
    try:
        ext = Path(filepath).suffix.lower()
        if ext == ".epub":
            return _convert_epub(Path(filepath))
        elif ext in {".html", ".htm"}:
            # Use the EPUB parsing pipeline to get the exact same heading processing
            from qmd.epub import (
                EPUBHTMLToMarkdown, clean_broken_paragraphs,
                promote_all_caps_headings, merge_consecutive_headings,
                normalize_headings
            )
            
            try:
                content = Path(filepath).read_text(encoding="utf-8")
            except UnicodeDecodeError:
                content = Path(filepath).read_text(encoding="latin-1", errors="replace")
                
            parser = EPUBHTMLToMarkdown(
                epub_zip=None,
                current_html_path=str(filepath),
                image_mode='refer'
            )
            parser.feed(content)
            md = parser.get_markdown()
            
            if md:
                # MOBI7 often fakes headings with bold/italic tags rather than h1-h6.
                # Promote standalone ***Text*** and **Text** to ATX headings before the pipeline.
                mobi_lines = []
                for line in md.splitlines():
                    s = line.strip()
                    if s.startswith('***') and s.endswith('***') and len(s) > 6:
                        mobi_lines.append(f"## {s[3:-3].strip()}")
                    elif s.startswith('**') and s.endswith('**') and len(s) > 4:
                        mobi_lines.append(f"### {s[2:-2].strip()}")
                    else:
                        mobi_lines.append(line)
                md = "\n".join(mobi_lines)

                md = clean_broken_paragraphs(md)
                md = promote_all_caps_headings(md, default_level=3)
                md = merge_consecutive_headings(md)
                md = normalize_headings(md)
                md = merge_consecutive_headings(md)
                md = re.sub(r'\n{3,}', '\n\n', md).strip()
            return md
        else:
            return _convert_text(Path(filepath))
    except Exception as e:
        raise ValueError(f"Failed to parse extracted MOBI file '{path.name}': {e}")
    finally:
        shutil.rmtree(tempdir, ignore_errors=True)


def _convert_html(path: Path) -> str:
    try:
        content = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        content = path.read_text(encoding="latin-1", errors="replace")

    try:
        from bs4 import BeautifulSoup
    except ImportError:
        raise ImportError("beautifulsoup4 is required for converting .html files. Install with `pip install beautifulsoup4`.")

    soup = BeautifulSoup(content, 'html.parser')

    for element in soup(["script", "style", "head", "title", "meta", "noscript", "svg", "iframe"]):
        element.decompose()

    md_lines: List[str] = []

    def process_node(node):
        if isinstance(node, str):
            text = node.strip()
            if text:
                md_lines.append(text)
            return

        name = node.name.lower() if node.name else ""

        if name in ["h1", "h2", "h3", "h4", "h5", "h6"]:
            level = int(name[1])
            prefix = "#" * level
            md_lines.append(f"\n{prefix} {node.get_text(strip=True)}\n")
        elif name == "p":
            text = node.get_text(strip=True)
            if text:
                md_lines.append(f"\n{text}\n")
        elif name in ["ul", "ol"]:
            for li in node.find_all("li", recursive=False):
                li_text = li.get_text(strip=True)
                if li_text:
                    md_lines.append(f"- {li_text}")
            md_lines.append("")
        elif name == "table":
            matrix = []
            for tr in node.find_all("tr"):
                row = [td.get_text(strip=True).replace('\n', ' ') for td in tr.find_all(["th", "td"])]
                if row:
                    matrix.append(row)
            table_md = _format_matrix_to_md_table(matrix)
            if table_md:
                md_lines.append(f"\n{table_md}\n")
        elif name in ["pre", "code"]:
            code_text = node.get_text()
            md_lines.append(f"\n```\n{code_text}\n```\n")
        else:
            for child in node.children:
                process_node(child)

    process_node(soup.body or soup)

    result = "\n".join(md_lines)
    result = re.sub(r'\n{3,}', '\n\n', result)
    return result.strip()


def guess_document_date(file_path: Union[str, Path], markdown_body: str = "") -> Optional[str]:
    """
    Attempts to infer the document date from the filename/path and markdown content using dplib.
    """
    dplib_mod = _get_dplib()
    if dplib_mod is None:
        return None
    try:
        sample_content = markdown_body[:4000] if markdown_body else ""
        if hasattr(dplib_mod, "extract_date"):
            try:
                res = dplib_mod.extract_date(path=str(file_path), content=sample_content)
            except TypeError:
                res = dplib_mod.extract_date(str(file_path), sample_content)
            if res:
                if hasattr(res, "hour") and res.hour == 0 and res.minute == 0 and res.second == 0 and res.microsecond == 0:
                    return res.strftime("%Y-%m-%d")
                return res.isoformat() if hasattr(res, "isoformat") else str(res)
        elif hasattr(dplib_mod, "DateResolver"):
            resolver = dplib_mod.DateResolver()
            report = resolver.resolve(filename=str(file_path).replace("\\", "/"), text_content=sample_content)
            if report and report.resolved_date:
                dt = report.resolved_date
                if dt.hour == 0 and dt.minute == 0 and dt.second == 0 and dt.microsecond == 0:
                    return dt.strftime("%Y-%m-%d")
                return dt.isoformat()
    except Exception:
        pass
    return None


def convert_to_markdown(file_path: Union[str, Path], config=None, errors_out: Optional[List[dict]] = None) -> str:
    """
    Converts a supported document or text file to Markdown.
    """
    path = Path(file_path)
    ext = path.suffix.lower()

    mod = sys.modules.get("qmd.converters")
    def _get_fn(name, fallback):
        return getattr(mod, name, fallback) if mod else fallback

    if ext in {".md", ".markdown", ".txt"}:
        raw_md = _get_fn("_convert_text", _convert_text)(path)
    elif ext == ".pdf":
        raw_md = _get_fn("_convert_pdf", _convert_pdf)(path, config, errors_out=errors_out)
    elif ext == ".docx":
        raw_md = _get_fn("_convert_docx", _convert_docx)(path, config, errors_out=errors_out)
    elif ext == ".pptx":
        raw_md = _get_fn("_convert_pptx", _convert_pptx)(path, config, errors_out=errors_out)
    elif ext == ".xlsx":
        raw_md = _get_fn("_convert_xlsx", _convert_xlsx)(path, config, errors_out=errors_out)
    elif ext == ".csv":
        raw_md = _get_fn("_convert_csv", _convert_csv)(path)
    elif ext in {".html", ".htm"}:
        raw_md = _get_fn("_convert_html", _convert_html)(path)
    elif ext == ".epub":
        raw_md = _get_fn("_convert_epub", _convert_epub)(path)
    elif ext == ".mobi":
        raw_md = _get_fn("_convert_mobi", _convert_mobi)(path)
    else:
        raw_md = _get_fn("_convert_text", _convert_text)(path)

    sanitized = _sanitize_text(raw_md)
    
    def _condense_table_block(match):
        block = match.group(0)
        if re.search(r'^[ \t]*\|[-\s:|]+\|[ \t]*$', block, flags=re.MULTILINE):
            return re.sub(r' {2,}', ' ', block)
        return block
        
    return re.sub(r'(?:^[ \t]*\|[^\n]*\|[ \t]*(?:\r?\n|$))+', _condense_table_block, sanitized, flags=re.MULTILINE)


def main():
    import argparse

    cur = Path(__file__).resolve().parent
    while cur.name != "src" and cur.parent != cur:
        cur = cur.parent
    src_dir = str(cur) if cur.name == "src" else str(Path(__file__).resolve().parents[2])
    if src_dir not in sys.path:
        sys.path.insert(0, src_dir)

    parser = argparse.ArgumentParser(
        description="Convert a document (.docx, .pptx, .xlsx, .csv, .html, .epub, .mobi, .md) to Markdown and inspect parsed blocks/chunks."
    )
    parser.add_argument("file_path", type=str, help="Path to the document to convert")
    parser.add_argument("-o", "--output", type=str, help="Optional output path to save the Markdown content")
    parser.add_argument("-v", "--verbose", action="store_true", help="Enable verbose debug output (shows inline raw responses from image endpoints)")
    parser.add_argument("--clean", action="store_true", help="Output only the raw converted Markdown without headers, stats, or metadata")
    parser.add_argument("--show-blocks", action="store_true", help="Display parsed semantic blocks from docparse")
    parser.add_argument("--show-chunks", action="store_true", help="Display chunked content ready for embedding")
    parser.add_argument("--vision-url", type=str, help="URL for the Vision API to test image extraction")
    parser.add_argument("--vision-api-key", type=str, help="Optional API key for the Vision API")
    parser.add_argument("--multimodal-url", type=str, help="URL for OpenAI-compatible multimodal endpoint")
    parser.add_argument("--multimodal-api-key", type=str, help="Optional API key for multimodal endpoint")
    parser.add_argument("--multimodal-model", type=str, help="Model name for multimodal endpoint")
    parser.add_argument("--max-image-concurrency", type=int, default=4, help="Max simultaneous images to process")

    args = parser.parse_args()

    file_path = Path(args.file_path).expanduser().resolve()
    if not file_path.exists():
        print(f"Error: File not found: {file_path}", file=sys.stderr)
        sys.exit(1)

    if not is_supported_file(file_path):
        print(f"Warning: Extension '{file_path.suffix}' is not explicitly supported. Attempting plain text conversion...", file=sys.stderr)

    if args.verbose:
        os.environ["QMD_VERBOSE"] = "1"

    mock_config = None
    if args.vision_url or args.multimodal_url or args.multimodal_model or args.verbose:
        class MockConfig:
            vision_url = args.vision_url
            vision_api_key = args.vision_api_key
            multimodal_url = args.multimodal_url
            multimodal_api_key = args.multimodal_api_key
            multimodal_model = args.multimodal_model
            max_image_concurrency = args.max_image_concurrency
            request_timeout = 120.0
            verbose = args.verbose
        mock_config = MockConfig()

    try:
        md_content = convert_to_markdown(file_path, config=mock_config)
    except Exception as e:
        print(f"Error converting {file_path}: {e}", file=sys.stderr)
        sys.exit(1)

    if args.clean:
        print(md_content)
    else:
        print("=" * 80)
        print(f"CONVERTED MARKDOWN ({file_path.name})")
        print("=" * 80)
        print(md_content)
        print("=" * 80)
        print(f"Stats: {len(md_content)} characters, {len(md_content.splitlines())} lines")
        inferred_date = guess_document_date(file_path, md_content)
        date_str = inferred_date if inferred_date else "None detected"
        print(f"Inferred Date: {date_str}")

    if args.output:
        out_path = Path(args.output).expanduser().resolve()
        out_path.write_text(md_content, encoding="utf-8")
        if not args.clean:
            print(f"\nSaved output to: {out_path}")

    if not args.clean and (args.show_blocks or args.show_chunks):
        try:
            from qmd.docparse.parser import parse_markdown_to_blocks
            from qmd.docparse.grouper import group_blocks_into_chunks

            blocks, has_tables = parse_markdown_to_blocks(content=md_content)
            print("\n" + "=" * 80)
            print(f"PARSED SEMANTIC BLOCKS ({len(blocks)} blocks found, has_tables={has_tables})")
            print("=" * 80)
            for i, block in enumerate(blocks, 1):
                print(f" Block #{i} [H{block.header_level}: {block.header_text}] (line {block.start_line}):")
                preview = block.content[:150].replace('\n', ' ')
                print(f"   {preview}..." if len(block.content) > 150 else f"   {preview}")

            if args.show_chunks:
                chunks = group_blocks_into_chunks(blocks, max_chunk_size=2048, target_chunk_size=1024)
                print("\n" + "=" * 80)
                print(f"GENERATED CHUNKS ({len(chunks)} chunks generated)")
                print("=" * 80)
                for i, chunk in enumerate(chunks, 1):
                    headers = " > ".join(chunk.parent_headers.values()) if chunk.parent_headers else "Root"
                    print(f"\n--- Chunk #{i} (Context: {headers}) ---")
                    print(chunk.content)
        except Exception as e:
            print(f"\nError running docparse analysis: {e}", file=sys.stderr)


if __name__ == "__main__":
    main()


__all__ = [
    "SUPPORTED_EXTENSIONS",
    "is_supported_file",
    "convert_to_markdown",
    "guess_document_date",
    "main",
    "_get_dplib",
    "_sanitize_text",
    "_format_matrix_to_md_table",
    "_convert_text",
    "_convert_pdf",
    "_convert_docx",
    "_convert_pptx",
    "_convert_xlsx",
    "_convert_csv",
    "_convert_epub",
    "_convert_mobi",
    "_convert_html",
    "_mathml_to_latex",
    "_omml_to_latex",
    "_extract_text_and_math",
    "_clean_vision_markdown",
    "clean_vision_text",
    "_wrap_vision_xml",
    "_is_image_processing_enabled",
    "_is_verbose",
    "_process_image_multimodal_llm",
    "_process_image",
    "_process_images_concurrently",
    "_process_image_vision_api",
    "build_header_detector",
    "clean_heading_text",
    "normalize_markdown_headings",
]